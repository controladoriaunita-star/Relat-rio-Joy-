"""Regressão contra a planilha: o motor precisa reproduzir os valores do Excel.

Lê as tabelas de dados, o DePara, a Projeção_2 e os saldos do .xlsm, calcula
Resumo e Painel com o motor e compara com os valores salvos pelo Excel em cada
célula. Toda célula de fórmula das colunas cobertas precisa ter equivalente.
"""
from __future__ import annotations

import math
from datetime import date, datetime, time

import openpyxl
import pytest
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

from backend.app.catalogo import VIABILIDADE_PADRAO, Premissas
from backend.app.config import ARQUIVO_EXCEL
from backend.app.excel_funcoes import DATA_ZERO
from backend.app.fontes.excel import ler_planilha
from backend.app.fontes.montagem import montar_base
from backend.app.excel_funcoes import chave_texto
from backend.app.motor.painel import calcular_painel
from backend.app.motor.resumo import calcular_resumo, receitas_do_mes

pytestmark = pytest.mark.skipif(not ARQUIVO_EXCEL.exists(), reason="planilha de referência ausente")

COLUNA_INICIO_TL = column_index_from_string("CE")
COLUNAS_RESUMO = {"viabilidade": "D", "pct_consumido": "F", "pago_ate": "H", "pago_mes": "J",
                  "a_pagar": "L", "projetado_mes": "Q", "projetado_proximo": "S", "total_mes": "U"}
COLUNAS_PAINEL = ("H", "L", "P", "T", "X", "Y", "AB")
COLUNAS_PAINEL_COBERTAS = set(COLUNAS_PAINEL) | {"J", "N", "BA", "BB", "FS", "FT", "FU", "FV"} | {
    get_column_letter(COLUNA_INICIO_TL + i) for i in range(84)}


@pytest.fixture(scope="module")
def cenario():
    planilha = ler_planilha(ARQUIVO_EXCEL)
    base, resultados = montar_base(planilha.tabelas, planilha.depara, planilha.projecao,
                                   planilha.ultima_atualizacao)
    hoje = planilha.ultima_atualizacao.date()
    # Z22 é digitado na planilha; no sistema é automático. Para comparar, usa o valor do Excel.
    premissas = Premissas(custo_projetado_proximo=775384.9700000001)
    formulas = openpyxl.load_workbook(ARQUIVO_EXCEL)
    return {
        "planilha": planilha, "resultados": resultados, "formulas": formulas, "base": base,
        "resumo": calcular_resumo(base, hoje, VIABILIDADE_PADRAO, premissas),
        "painel": calcular_painel(base, hoje, VIABILIDADE_PADRAO, premissas),
    }


def _eh_formula(valor) -> bool:
    return isinstance(valor, ArrayFormula) or (isinstance(valor, str) and valor.startswith("="))


def _normalizar(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, time):          # data vazia formatada (0)
        return 0.0
    return valor


def _igual(esperado, obtido) -> bool:
    esperado, obtido = _normalizar(esperado), _normalizar(obtido)
    if isinstance(obtido, date) and isinstance(esperado, (int, float)):   # data em formato número
        obtido = float((obtido - DATA_ZERO).days)
    if isinstance(esperado, (int, float)) and not isinstance(esperado, bool):
        if not isinstance(obtido, (int, float)):
            return esperado == 0 and obtido is None
        return math.isclose(esperado, obtido, rel_tol=1e-9, abs_tol=0.005)
    if esperado in (None, ""):
        return obtido in (None, "", 0)
    return esperado == obtido


def _comparar(valores_excel, celulas: dict, rotulo: str, ajustes: dict | None = None) -> None:
    """`ajustes` = célula → diferença esperada (motor − Excel) por correção aprovada."""
    ajustes = ajustes or {}
    esperado = lambda c: valores_excel[c].value + ajustes[c] if c in ajustes else valores_excel[c].value
    divergencias = [f"{rotulo}!{c}: Excel={valores_excel[c].value!r} motor={v!r}"
                    for c, v in celulas.items() if not _igual(esperado(c), v)]
    assert not divergencias, f"{len(divergencias)} divergência(s):\n" + "\n".join(divergencias[:40])


def _ajuste_saidas(cenario, resumo) -> dict:
    """Correção aprovada: Controle Saídas entram pelo mês da data. A planilha só contava como
    "pago no mês" (J) a saída do último dia do mês; as outras do mês caíam em "pago até" (H).
    O valor dessas saídas passa de H para J e entra no custo do mês e no aporte."""
    mes = resumo["datas"]["mes_atual"]
    ajustes = {}
    for linha in resumo["linhas"]:
        delta = sum(m.valor for m in cenario["base"].filtrar("saidas", "classificacao", linha["rubrica"])
                    if m.fim_mes == mes and m.data != mes)
        if delta:
            r = linha["linha_excel"]
            for celula, sinal in ((f"H{r}", -1), (f"J{r}", 1), (f"U{r}", 1), ("H5", -1), ("J5", 1),
                                  ("U5", 1), ("Z7", 1), ("Z9", 1)):
                ajustes[celula] = ajustes.get(celula, 0) + sinal * delta
    return ajustes


def _ajuste_receitas(cenario, resumo) -> dict:
    """Correção aprovada (out/2026): a receita do mês (Z6) é tudo o que a aba Recebidas registra
    no mês. A planilha usava as entradas do Controle Entradas que não são aporte (rendimentos).
    A diferença entra com sinal trocado no aporte (Z9)."""
    mes = resumo["datas"]["mes_atual"]
    aportes = {chave_texto(c) for c in ("Aporte", "Aporte Holding")}
    planilha = sum(m.valor for m in cenario["base"].entradas if m.fim_mes == mes
                   and chave_texto((m.classificacao or "").strip()) not in aportes)
    delta = receitas_do_mes(cenario["base"], mes) - planilha
    return {"Z6": delta, "Z9": -delta}


def _ajustes_resumo(cenario, resumo) -> dict:
    ajustes = _ajuste_saidas(cenario, resumo)
    for celula, delta in _ajuste_receitas(cenario, resumo).items():
        ajustes[celula] = ajustes.get(celula, 0) + delta
    if "Z9" in ajustes:
        ajustes["Z10"] = ajustes["Z9"] * 0.85
        ajustes["Z11"] = ajustes["Z9"] * 0.15
    return ajustes


def test_classificacao_depara_igual_a_planilha(cenario):
    """O DePara aplicado pelo sistema gera a mesma Classificação que a planilha."""
    for codigo, resultado in cenario["resultados"].items():
        assert not resultado.colunas_ausentes, codigo
        assert not resultado.pendentes, (codigo, resultado.pendentes[:3])


def test_resumo(cenario):
    resumo = cenario["resumo"]
    celulas = {}
    for linha in resumo["linhas"] + [{**resumo["total"], "linha_excel": 5}]:
        for campo, coluna in COLUNAS_RESUMO.items():
            if campo in linha:
                celulas[f"{coluna}{linha['linha_excel']}"] = linha[campo]
    blocos = {"aporte_mes": [("saldo_anterior", 5), ("receita_mes", 6), ("custo_mes", 7),
                             ("caixa_minimo", 8), ("aporte_total", 9), ("investidor", 10), ("joy", 11)],
              "aporte_proximo": [("receita_projetada", 21), ("custo_projetado", 22), ("caixa_minimo", 23)]}
    for bloco, campos in blocos.items():
        for campo, linha in campos:
            celulas[f"Z{linha}"] = resumo[bloco][campo]
    # Y32:Y43 (aportes programados) viraram o cronograma do aporte do mês: as parcelas somam
    # o aporte necessário (Z9) em vez de valores fixos digitados. Coberto em test_regras_aporte.
    titulos = resumo["titulos"]
    celulas.update({"H3": titulos["pago_ate"], "J3": titulos["pago_mes"], "L3": titulos["a_pagar"],
                    "S3": titulos["projetado_proximo"], "Y5": titulos["saldo_anterior"],
                    "B4": resumo["datas"]["mes_atual"],
                    "L4": resumo["datas"]["a_pagar_ate"], "AD3": resumo["datas"]["mes_analise"]})
    aporte = resumo["aporte_mes"]["aporte_total"]
    assert aporte > 0, "o ajuste de Z9 supõe aporte positivo (sem o piso de zero)"
    _comparar(cenario["planilha"].livro["Resumo"], celulas, "Resumo", _ajustes_resumo(cenario, resumo))

    # Toda fórmula das colunas de valor precisa estar coberta.
    formulas = cenario["formulas"]["Resumo"]
    faltando = [c.coordinate for linha in formulas.iter_rows(min_row=5, max_row=43)
                for c in linha if c.column_letter in COLUNAS_RESUMO.values()
                and _eh_formula(c.value) and c.coordinate not in celulas]
    assert not faltando, faltando


# Correções aprovadas na revisão do Painel (out/2026; ver motor/painel.py, P1..P11). As células
# abaixo podem divergir da planilha; todas as outras precisam bater. As regras novas são
# verificadas em test_regras_aporte.py (seção Painel) e nos testes de invariantes abaixo.
COLUNAS_CORRIGIDAS = {"L", "P", "T", "AB", "FS", "FT", "FU", "FV"}          # P2, P3, P4, P10
LINHAS_TOTAIS_CORRIGIDAS = {59, 79, 83, 93, 95, 101}                       # P6, P11, P4
LINHAS_TL_CORRIGIDAS = {32, 36, 41, 105, 106, 119, 122}                    # P5, P7, P9
LINHAS_SEM_EVE = set(range(105, 117)) | {119}                              # P11
LINHAS_REMOVIDAS = {117}                                                   # auxiliar sem uso


def _corrigida(celula: str, mes_analise, meses) -> bool:
    coluna = celula.rstrip("0123456789")
    linha = int(celula[len(coluna):])
    if linha in LINHAS_REMOVIDAS or (linha == 122 and coluna in ("J", "N")):
        return True
    if coluna in COLUNAS_CORRIGIDAS:
        return True
    if coluna in ("H", "X", "Y") and linha in LINHAS_TOTAIS_CORRIGIDAS:
        return True
    if coluna == "X" and linha in LINHAS_SEM_EVE:
        return True
    i = column_index_from_string(coluna) - COLUNA_INICIO_TL
    if 0 <= i < len(meses) and linha >= 24:
        if linha in LINHAS_TL_CORRIGIDAS or linha in LINHAS_TOTAIS_CORRIGIDAS:
            return True
        if 43 <= linha <= 76 and meses[i] <= mes_analise:                  # P1: só realizado
            return True
    return False


def test_painel(cenario):
    painel = cenario["painel"]
    celulas = {}
    for linha in painel["linhas"]:
        r = linha.linha_excel
        for coluna in COLUNAS_PAINEL:
            valor = getattr(linha, coluna)
            if valor is not None:
                celulas[f"{coluna}{r}"] = valor
        for i, valor in enumerate(linha.tl):
            if valor is not None:
                celulas[f"{get_column_letter(COLUNA_INICIO_TL + i)}{r}"] = valor
        for coluna, valor in linha.extras.items():
            celulas[f"{coluna}{r}"] = valor
    for i, mes in enumerate(painel["meses"]):
        coluna = get_column_letter(COLUNA_INICIO_TL + i)
        celulas[f"{coluna}4"] = celulas[f"{coluna}23"] = celulas[f"{coluna}91"] = mes
    celulas.update({"D22": painel["datas"]["mes_analise"], "H21": painel["datas"]["mes_atual"],
                    "AI23": painel["datas"]["corte"], "G20": painel["titulos"]["ate"],
                    "H23": painel["titulos"]["realizado"]})

    valores = cenario["planilha"].livro["Painel"]
    formulas = cenario["formulas"]["Painel"]
    corrigida = lambda c: _corrigida(c, painel["datas"]["mes_analise"], painel["meses"])
    # Só compara células que existem na planilha (as demais são vazias lá) e não corrigidas.
    existentes = {c: v for c, v in celulas.items() if formulas[c].value is not None and not corrigida(c)}
    assert len(existentes) > 2500, len(existentes)          # a maior parte do Painel segue a planilha
    _comparar(valores, existentes, "Painel")

    faltando = [c.coordinate for linha in formulas.iter_rows(min_row=24, max_row=122)
                for c in linha if _eh_formula(c.value)
                and get_column_letter(c.column) in COLUNAS_PAINEL_COBERTAS
                and c.coordinate not in celulas and not corrigida(c.coordinate)]
    assert not faltando, faltando


def test_painel_invariantes_com_dados_da_planilha(cenario):
    """Com os dados reais da planilha, as regras da revisão se sustentam."""
    painel = cenario["painel"]
    linhas = {l.linha_excel: l for l in painel["linhas"]}
    for l in painel["linhas"]:                               # linha do tempo × colunas
        if "FT" in l.extras:
            assert l.extras["FT"] == l.extras["FU"] == "OK!", l.linha_excel
    i = painel["meses"].index(painel["datas"]["mes_analise"])
    assert math.isclose(linhas[122].tl[i], linhas[119].H, abs_tol=0.005)    # acumulado = saldo A
    # Realizado de custos e despesas (A) continua igual à planilha.
    excel = cenario["planilha"].livro["Painel"]
    for r in range(43, 77):
        if r in linhas and linhas[r].grupo in ("custo", "despesa") and excel[f"H{r}"].value is not None:
            assert math.isclose(linhas[r].H, excel[f"H{r}"].value, abs_tol=0.005), r
