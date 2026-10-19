"""Leitura do Controle - Joy+ Evans.xlsm.

Serve para a carga inicial do sistema (DePara, saldos, projeção) e como base do
teste de regressão: os valores calculados pelo motor são comparados com os
valores que o próprio Excel salvou nas abas Resumo e Painel.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import openpyxl

from ..excel_funcoes import como_data, como_numero
from ..motor.base import LinhaProjecao, SaldoConta
from .definicoes import FONTES


@dataclass
class PlanilhaExcel:
    tabelas: dict[str, list[list]] = field(default_factory=dict)     # código da fonte → valores
    depara: dict[str, list[tuple]] = field(default_factory=dict)     # código → [(concat, classif)]
    projecao: list[LinhaProjecao] = field(default_factory=list)
    saldos: list[SaldoConta] = field(default_factory=list)
    ultima_atualizacao: datetime | None = None
    livro: object = None                                              # workbook (valores salvos)


def _valores_tabela(livro, nome_aba: str, nome_tabela: str) -> list[list]:
    aba = livro[nome_aba]
    intervalo = aba.tables[nome_tabela].ref
    return [[c.value for c in linha] for linha in aba[intervalo]]


def _tabela_por_nome(livro, nome_tabela: str) -> list[list]:
    for aba in livro.worksheets:
        if nome_tabela in aba.tables:
            return _valores_tabela(livro, aba.title, nome_tabela)
    raise KeyError(nome_tabela)


def ler_planilha(caminho: Path) -> PlanilhaExcel:
    livro = openpyxl.load_workbook(caminho, data_only=True)
    planilha = PlanilhaExcel(livro=livro)

    for fonte in FONTES:
        planilha.tabelas[fonte.codigo] = _valores_tabela(livro, fonte.aba_excel, fonte.tabela_excel)
        if fonte.tabela_depara:
            valores = _tabela_por_nome(livro, fonte.tabela_depara)
            planilha.depara[fonte.codigo] = [(linha[1], linha[0]) for linha in valores[1:]]

    cab, *linhas = _tabela_por_nome(livro, "Projeção_2")
    i = {nome: cab.index(nome) for nome in ("Descrição", "Datas", "Tipo", "Valor")}
    planilha.projecao = [LinhaProjecao(l[i["Descrição"]], como_data(l[i["Datas"]]),
                                       l[i["Tipo"]], como_numero(l[i["Valor"]])) for l in linhas]

    cab, *linhas = _tabela_por_nome(livro, "Consulta2")
    i = {nome: cab.index(nome) for nome in ("Empresa", "Conta", "Data", "Saldo")}
    planilha.saldos = [SaldoConta(l[i["Empresa"]], l[i["Conta"]], como_data(l[i["Data"]]),
                                  como_numero(l[i["Saldo"]]))
                       for l in linhas if l[i["Data"]] is not None]

    _, *linhas = _tabela_por_nome(livro, "Ultima_Atualização")
    planilha.ultima_atualizacao = linhas[0][0] if linhas else None
    return planilha
